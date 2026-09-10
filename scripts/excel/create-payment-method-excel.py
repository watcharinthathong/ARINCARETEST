"""
create-payment-method-excel.py
สร้าง Excel test case report สำหรับ retest ฟีเจอร์ "เพิ่มวิธีการชำระเงินแบบกำหนดเอง" (POS)
Sheets: Test Cases, Bug Log, สรุป & Coverage
"""

import os
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(BASE, "docs", "test-cases", "Arincare_POS_PaymentMethod_TestCases.xlsx")

# ── Styles (ตาม convention เดิมของโปรเจกต์: docs/test-cases/Arincare_CustomerPOS_TestCases.xlsx) ──
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Tahoma", bold=True, color="FFFFFF", size=11)
BODY_FONT   = Font(name="Tahoma", size=10)
BOLD_FONT   = Font(name="Tahoma", size=10, bold=True)
PASS_FILL   = PatternFill("solid", fgColor="C6EFCE")
PASS_FONT   = Font(name="Tahoma", bold=True, color="276221", size=10)
FAIL_FILL   = PatternFill("solid", fgColor="FFC7CE")
FAIL_FONT   = Font(name="Tahoma", bold=True, color="9C0006", size=10)
SKIP_FILL   = PatternFill("solid", fgColor="FFEB9C")
SKIP_FONT   = Font(name="Tahoma", bold=True, color="9C6500", size=10)
WRAP_TOP    = Alignment(wrap_text=True, vertical="top")
CENTER      = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin        = Side(style="thin", color="CCCCCC")
BORDER      = Border(left=thin, right=thin, top=thin, bottom=thin)

SS = "screenshots/payment-method-retest/"

# ── Test Cases ───────────────────────────────────────────────────────────────
# (tc_id, tab, priority, type, precondition, test_data, steps, expected, actual, status, notes)
CASES = [
    (
        "TC-PAY-001", "POS - ตะกร้าสินค้า", "High", "Positive (Happy)",
        "Login POS สำเร็จ บริษัท Watcharin TestTest สาขา สำนักงานใหญ่",
        "สินค้า \"BILAXTEN 20 MG TABLETS 10'S\" หน่วย \"แผง\"",
        "1. ค้นหาสินค้าด้วยช่องค้นหา (ctrl+Q)\n2. เลือกสินค้าจากผลค้นหา\n3. เลือกหน่วย \"แผง\" (หน่วย \"ขวด\" default ราคา 0 บาท ขายไม่ได้)\n4. กดเพิ่มลงตะกร้า",
        "สินค้าถูกเพิ่มลงตะกร้า แสดงราคาถูกต้อง (110.00 บาท)",
        "เพิ่มลงตะกร้าสำเร็จ ราคา 110.00 บาทแสดงถูกต้องตรงตามหน่วย \"แผง\"",
        "Pass",
        SS + "01-pos-main-after-login.png ... 06-after-add-to-cart.png",
    ),
    (
        "TC-PAY-002", "POS - Checkout", "High", "Positive (Happy)",
        "มีสินค้าในตะกร้าอย่างน้อย 1 รายการ (TC-PAY-001)",
        "-",
        "1. กดปุ่ม \"ชำระเงิน\"\n2. ผ่าน modal คำเตือนยาอันตราย (ถ้ามี)\n3. ผ่าน modal Medical Adherence (ถ้ามี)",
        "เข้าสู่หน้าสรุปยอด (checkout summary) พร้อมยอดเงินรวมถูกต้อง",
        "เข้าหน้าสรุปยอดสำเร็จ ผ่าน 2 modal คั่นกลางถูกต้อง ยอดเงินรวมแสดง 110.00 บาท",
        "Pass",
        SS + "07-checkout-summary-or-warning.png, 08-checkout-summary.png",
    ),
    (
        "TC-PAY-003", "POS - Payment Method", "Critical", "Positive (Happy)",
        "อยู่หน้าสรุปยอด (checkout summary)",
        "-",
        "1. กดปุ่ม \"+\" มุมล่างแถวไอคอนวิธีการชำระเงิน (เปิด \"จัดการวิธีการชำระเงิน\")\n2. กดปุ่ม \"+\" ชั้นในอีกครั้ง (Add New Payment)",
        "เปิดฟอร์ม \"เพิ่ม Payment Method\" ให้กรอกชื่อ + อัปโหลดรูปไอคอน",
        "ฟอร์ม \"เพิ่ม Payment Method\" เปิดถูกต้อง มีช่องกรอกชื่อและปุ่มอัปโหลดรูป",
        "Pass",
        SS + "09-add-payment-modal.png, 10-new-payment-method-form.png",
    ),
    (
        "TC-PAY-004", "POS - Payment Method", "Critical", "Positive (Happy)",
        "อยู่ที่ฟอร์ม \"เพิ่ม Payment Method\"",
        "ชื่อวิธีชำระ = QA<8 หลักท้าย timestamp> (<=15 ตัวอักษร), รูปจาก ~/Downloads/010011598.jpg",
        "1. กรอกชื่อวิธีการชำระเงิน\n2. อัปโหลดรูปจาก Downloads\n3. กด \"Crop Photo\"\n4. กด \"ยืนยัน\"",
        "สร้างวิธีการชำระเงินใหม่สำเร็จ กลับหน้าสรุปยอดอัตโนมัติ และวิธีชำระใหม่ปรากฏในรายการช่องทางชำระเงินจริง",
        "สร้างสำเร็จ กลับหน้าสรุปยอดอัตโนมัติ วิธีชำระใหม่ปรากฏในรายการไอคอนช่องทางชำระเงินจริง",
        "Pass",
        SS + "11-new-payment-form-filled.png, 12-new-payment-form-cropped.png, 13-after-confirm-new-payment-method.png, 14-checkout-summary-after-new-method.png",
    ),
    (
        "TC-PAY-005", "POS - Payment Method", "High", "Positive (Happy)",
        "วิธีการชำระเงินใหม่ถูกสร้างแล้ว (TC-PAY-004)",
        "-",
        "1. เลือกวิธีการชำระเงินที่เพิ่งสร้าง\n2. สังเกตช่องหมายเหตุที่ปรากฏ\n3. ลองกดชำระเงินโดยไม่กรอกหมายเหตุ",
        "มีช่องหมายเหตุ (remark) ปรากฏให้กรอก และต้องเป็น optional (ปุ่มชำระเงินไม่ถูก disable แม้ไม่กรอกหมายเหตุ)",
        "ช่องหมายเหตุปรากฏถูกต้อง ยืนยันเป็น optional จริง — ปุ่มชำระเงินกดได้แม้ไม่กรอกข้อความ",
        "Pass",
        SS + "15-new-method-selected.png",
    ),
    (
        "TC-PAY-006", "POS - Payment Method", "Critical", "Positive (Happy)",
        "เลือกวิธีการชำระเงินใหม่แล้ว อยู่ระหว่างกรอกหมายเหตุ",
        "หมายเหตุ = \"QA-TEST-หมายเหตุ-<timestamp>\"",
        "1. กรอกข้อความหมายเหตุเฉพาะเจาะจง\n2. กดปุ่ม \"ชำระด้วย <ชื่อวิธีชำระ>\"\n3. กด \"ทำรายการต่อไป\" เมื่อชำระสำเร็จ",
        "ชำระเงินสำเร็จ ระบบบันทึกรายการขายพร้อมหมายเหตุที่กรอกไว้",
        "ชำระเงินสำเร็จ ขึ้นข้อความ \"บันทึกรายการขายเสร็จสิ้น\"",
        "Pass",
        SS + "16-remark-filled.png, 17-after-complete-sale.png",
    ),
    (
        "TC-PAY-007", "POS - รายการบิล", "Medium", "UI",
        "มีบิลที่ชำระด้วยวิธีการชำระเงินแบบกำหนดเองพร้อมหมายเหตุ (TC-PAY-006)",
        "-",
        "1. เปิดหน้า \"รายการบิล\" ในฝั่ง POS\n2. หาบิลที่เพิ่งสร้าง\n3. Hover/คลิกไอคอน info (ⓘ) ที่คอลัมน์ \"ประเภทการชำระเงิน\"",
        "Tooltip แสดงข้อความหมายเหตุที่กรอกไว้ตอนชำระเงิน ตรงกับที่กรอกจริง",
        "ยืนยันด้วยภาพจริงจากผู้ใช้ (เบราว์เซอร์จริง ไม่ใช่ automation): tooltip แสดงข้อความ \"QA-TEST-หมายเหตุ-...\" ถูกต้องตรงกับที่กรอกไว้ 100%. "
        "หมายเหตุอัตโนมัติ: Playwright hover/click ไม่สามารถ trigger tooltip นี้ได้เลย (สืบสวนละเอียดหลายวิธี — ดู docs/payment-method-discovery/verify-tooltip-fullpage-bug.ts) "
        "จึง skip TC ฝั่ง automated regression ไว้ก่อน แต่ manual QA ยืนยัน Pass ชัดเจน",
        "Pass",
        SS + "18-pos-bill-list.png, 19-bill-list-tooltip-hover.png (ยืนยันเพิ่มเติมจากภาพที่ผู้ใช้ส่งเอง)",
    ),
    (
        "TC-PAY-008", "ERP - list-bills", "Medium", "UI",
        "Login app-stg.arincare.com บริษัท Watcharin TestTest สาขาที่ถูกต้อง, มีบิลจาก TC-PAY-006",
        "URL: https://app-stg.arincare.com/companies/list-bills",
        "1. เข้าหน้า \"รายการบิลขาย\"\n2. ค้นหา/หาบิลที่เพิ่งสร้าง\n3. Hover/คลิกไอคอน info (ⓘ) ที่คอลัมน์ \"ประเภทการชำระเงิน\"",
        "Tooltip แสดงข้อความหมายเหตุตรงกับที่กรอกไว้ตอนชำระเงินฝั่ง POS (ต้องตรงกับ TC-PAY-007)",
        "ยืนยันด้วยภาพจริงจากผู้ใช้: tooltip แสดงข้อความหมายเหตุตรงกับฝั่ง POS 100% (บิลเดียวกัน QA35016112). "
        "เช่นเดียวกับ TC-PAY-007 — Playwright automation ไม่สามารถ trigger ได้ แต่ manual QA ยืนยัน Pass",
        "Pass",
        SS + "106-erp-list-bills-searched.png, 107-erp-payment-tooltip-hover.png",
    ),
    (
        "TC-PAY-009", "ERP - list-bills", "High", "Integration",
        "อยู่หน้า \"รายการบิลขาย\" (list-bills) มีบิลจาก TC-PAY-006",
        "-",
        "1. หาบิลที่เพิ่งสร้าง\n2. กดปุ่ม \"รายละเอียดสินค้า\" (คอลัมน์ขวาสุด)\n3. ตรวจสอบรายการสินค้า/ราคา/จำนวนในหน้าต่างรายละเอียด",
        "ข้อมูลสินค้า (BILAXTEN, จำนวน 1 แผง, ราคา 110.00 บาท) ตรงกับที่ขายจริงในหน้า POS",
        "ข้อมูลสินค้า/ราคา/จำนวนตรงกับที่ขายจริงทุกรายการ (BILAXTEN 1 แผง 110.00 บาท)",
        "Pass",
        SS + "109-erp-product-detail-modal.png",
    ),
    (
        "TC-PAY-010", "ERP - รายงานกำไร-ขาดทุน/บิล", "Medium", "Integration",
        "มีบิลจาก TC-PAY-006 ปรากฏในระบบแล้ว",
        "URL: https://app-stg.arincare.com/companies/reports/sales?report=profitBillSummary",
        "1. เข้าหน้ารายงานกำไร-ขาดทุน/บิล\n2. ค้นหาบิลที่เพิ่งสร้าง\n3. ตรวจสอบยอดขาย, ภาษี, วิธีการชำระเงินในรายงาน เทียบกับบิลจริง",
        "ตัวเลขยอดขาย/ภาษี และวิธีการชำระเงินที่แสดงในรายงานตรงกับบิลจริง (ยอดขาย 110.00, ภาษี 7.20 บาท)",
        "ตัวเลขยอดขาย 110.00 บาท ภาษี 7.20 บาท และวิธีการชำระเงินตรงกับบิลจริงทั้งหมด",
        "Pass",
        SS + "111-erp-profit-report-searched.png",
    ),
]

# ── Bug Log (เก็บไว้เพื่อ traceability แม้ retract แล้ว) ────────────────────────
# (bug_id, severity, title, status, description, linked_tc)
BUGS = [
    (
        "BUG-001", "N/A (Retracted)",
        "[RETRACTED] ไอคอน info ที่คอลัมน์ประเภทการชำระเงินไม่แสดง tooltip หมายเหตุ",
        "Retracted — False Positive",
        "พบระหว่าง exploratory retest รอบแรก: Playwright hover/click ไอคอน info (ⓘ) แล้ว tooltip ไม่ขึ้นเลย "
        "ทั้งฝั่ง POS และ ERP list-bills — ตอนแรกสรุปว่าเป็น Major bug\n\n"
        "แก้ไข: ผู้ใช้ส่ง screenshot จากเบราว์เซอร์จริง (ไม่ใช่ automation) ยืนยันว่า tooltip แสดงข้อความหมายเหตุถูกต้อง "
        "100% บนบิลเดียวกันเป๊ะ (QA35073759, QA35016112) ทั้งฝั่ง POS และ ERP — ฟีเจอร์ทำงานถูกต้องจริง\n\n"
        "สืบสวนซ้ำด้วย Playwright หลายวิธี (hover ปกติ, click, mousedown/up จริง, mouse.move แบบ steps, "
        "headed mode, ปิด navigator.webdriver) — DOM/network ไม่ขยับเลยทุกครั้ง สาเหตุที่ automation "
        "trigger ไม่ได้ยังไม่ทราบแน่ชัด (icon ไม่มี title/data-xxx/aria-xxx attribute ใน static DOM เลย) "
        "แต่ยืนยันแล้วว่าไม่ใช่ product bug — เป็นข้อจำกัดของการทดสอบอัตโนมัติเท่านั้น",
        "TC-PAY-007, TC-PAY-008",
    ),
]

# ── Build workbook ───────────────────────────────────────────────────────────
wb = openpyxl.Workbook()

# --- Sheet 1: Test Cases ---
ws = wb.active
ws.title = "Test Cases"
headers = ["TC-ID", "Tab/หน้าจอ", "Priority", "Type", "Precondition",
           "Test Data", "Steps", "Expected Result", "Actual Result", "Status", "Notes"]
widths = [12, 22, 10, 16, 28, 26, 34, 34, 40, 11, 34]

for col, (h, w) in enumerate(zip(headers, widths), start=1):
    c = ws.cell(1, col, h)
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = CENTER
    c.border = BORDER
    ws.column_dimensions[get_column_letter(col)].width = w
ws.freeze_panes = "A2"

STATUS_STYLE = {
    "Pass": (PASS_FILL, PASS_FONT),
    "Fail": (FAIL_FILL, FAIL_FONT),
    "Skipped": (SKIP_FILL, SKIP_FONT),
    "Blocked": (SKIP_FILL, SKIP_FONT),
}

for i, row in enumerate(CASES, start=2):
    for col, val in enumerate(row, start=1):
        c = ws.cell(i, col, val)
        c.border = BORDER
        c.font = BOLD_FONT if col == 1 else BODY_FONT
        c.alignment = CENTER if col in (1, 3, 4, 10) else WRAP_TOP
    status_cell = ws.cell(i, 10)
    fill, font = STATUS_STYLE.get(status_cell.value, (None, BODY_FONT))
    if fill:
        status_cell.fill = fill
        status_cell.font = font
    ws.row_dimensions[i].height = 90

# --- Sheet 2: Bug Log ---
ws2 = wb.create_sheet("Bug Log")
bug_headers = ["Bug ID", "Severity", "Title", "Status", "Description", "Linked TC"]
bug_widths = [12, 18, 40, 22, 70, 20]
for col, (h, w) in enumerate(zip(bug_headers, bug_widths), start=1):
    c = ws2.cell(1, col, h)
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = CENTER
    c.border = BORDER
    ws2.column_dimensions[get_column_letter(col)].width = w
ws2.freeze_panes = "A2"

for i, row in enumerate(BUGS, start=2):
    for col, val in enumerate(row, start=1):
        c = ws2.cell(i, col, val)
        c.border = BORDER
        c.font = BOLD_FONT if col == 1 else BODY_FONT
        c.alignment = CENTER if col in (1, 2, 4, 6) else WRAP_TOP
    ws2.cell(i, 4).fill = SKIP_FILL
    ws2.cell(i, 4).font = SKIP_FONT
    ws2.row_dimensions[i].height = 200

# --- Sheet 3: สรุป & Coverage ---
ws3 = wb.create_sheet("สรุป & Coverage")
ws3.column_dimensions["A"].width = 45
ws3.column_dimensions["B"].width = 14

def section_title(row, text):
    c = ws3.cell(row, 1, text)
    c.font = Font(name="Tahoma", bold=True, size=12, color="1F4E78")
    return row + 1

r = 1
r = section_title(r, "QA Sign-off Summary — ฟีเจอร์ เพิ่มวิธีการชำระเงินแบบกำหนดเอง (POS)")
r += 1

meta = [
    ("Environment", "Staging (pos-stg.arincare.com + app-stg.arincare.com)"),
    ("Company / Branch", "Watcharin TestTest / สำนักงานใหญ่"),
    ("วันที่ทดสอบ", "2026-09-02"),
    ("จำนวน Test Case ทั้งหมด", len(CASES)),
    ("Pass", sum(1 for c in CASES if c[9] == "Pass")),
    ("Fail", sum(1 for c in CASES if c[9] == "Fail")),
    ("Bug ที่พบ (Active)", 0),
    ("Bug ที่ Retract แล้ว (False Positive)", len(BUGS)),
    ("Go/No-Go Decision", "GO"),
]
for k, v in meta:
    ws3.cell(r, 1, k).font = BOLD_FONT
    ws3.cell(r, 2, v).font = BODY_FONT
    ws3.cell(r, 1).border = BORDER
    ws3.cell(r, 2).border = BORDER
    r += 1

r += 1
r = section_title(r, "สรุปตาม Priority")
ws3.cell(r, 1, "Priority").font = HEADER_FONT
ws3.cell(r, 1).fill = HEADER_FILL
ws3.cell(r, 2, "จำนวน").font = HEADER_FONT
ws3.cell(r, 2).fill = HEADER_FILL
r += 1
from collections import Counter
pri_counts = Counter(c[2] for c in CASES)
for pri in ["Critical", "High", "Medium", "Low"]:
    if pri in pri_counts:
        ws3.cell(r, 1, pri).font = BODY_FONT
        ws3.cell(r, 2, pri_counts[pri]).font = BODY_FONT
        r += 1
ws3.cell(r, 1, "รวม").font = BOLD_FONT
ws3.cell(r, 2, len(CASES)).font = BOLD_FONT
r += 2

r = section_title(r, "สรุปตามประเภท (Type)")
ws3.cell(r, 1, "Type").font = HEADER_FONT
ws3.cell(r, 1).fill = HEADER_FILL
ws3.cell(r, 2, "จำนวน").font = HEADER_FONT
ws3.cell(r, 2).fill = HEADER_FILL
r += 1
type_counts = Counter(c[3] for c in CASES)
for t, n in sorted(type_counts.items()):
    ws3.cell(r, 1, t).font = BODY_FONT
    ws3.cell(r, 2, n).font = BODY_FONT
    r += 1
r += 1

r = section_title(r, "Coverage — สิ่งที่ครอบคลุม")
coverage_notes = [
    "เพิ่มสินค้าลงตะกร้า + เลือกหน่วยที่มีราคาถูกต้อง",
    "สร้างวิธีการชำระเงินแบบกำหนดเอง (ชื่อ + อัปโหลดรูป + Crop)",
    "เลือกใช้วิธีการชำระเงินที่สร้างใหม่ + ทดสอบว่าช่องหมายเหตุเป็น optional จริง",
    "ชำระเงินพร้อมหมายเหตุ end-to-end",
    "Cross-verification ข้าม 3 หน้าจอ: POS รายการบิล / ERP list-bills / ERP รายงานกำไร-ขาดทุน",
    "ตรวจสอบความถูกต้องของตัวเลข (ยอดขาย, ภาษี) และรายละเอียดสินค้าข้ามหน้าจอ",
    "Manual verification ยืนยันผลจากภาพจริงของผู้ใช้ เมื่อ automation ให้ผลขัดแย้งกับความเป็นจริง",
]
for note in coverage_notes:
    ws3.cell(r, 1, "• " + note).font = BODY_FONT
    ws3.cell(r, 1).alignment = WRAP_TOP
    r += 1

wb.save(OUT)
print(f"Saved: {OUT}")
print(f"Test Cases: {len(CASES)} | Bugs logged: {len(BUGS)}")
